from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, text in enumerate(headers):
        header_cells[index].text = text

    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = text

    doc.add_paragraph("")


def build_document() -> Document:
    doc = Document()

    doc.add_heading("以下为前端需要展示的全部数据汇总表", level=0)
    doc.add_paragraph("更新日期：2026-07-08。")
    doc.add_paragraph("本版已按当前前端实际实现更新，涵盖全局状态条、事件时间线、车辆控制历史、联锁冲突检查、供电扩展分析、轨道筛选等新增内容。")
    doc.add_paragraph("一、数据总览")
    doc.add_paragraph("主 UI 当前统一走单条 WebSocket（/ws/dashboard）+ 静态线路 line-layout.json；控车走 POST /api/v1/vehicle/control；其余展示与分析项由前端 store 在实时快照基础上继续计算。")

    add_table(
        doc,
        ["数据来源", "加载方式", "用途"],
        [
            ["静态线路 line-layout.json", "GET /data/line-layout.json", "几何、站名、区段边界、信号位置、道岔静态关系、拓扑图、坡度"],
            ["实时快照 dashboard_snapshot", "WebSocket /ws/dashboard", "列车、区段占用、信号状态、道岔状态、供电、告警、系统信息"],
            ["控车指令", "POST /api/v1/vehicle/control", "手动牵引 / 制动 / 紧急制动，驱动车辆页按钮、键盘控车与控制历史"],
            ["前端自算", "Simulation / UI Store 内计算", "历史曲线、占线时间轴、停车记录、事件时间线、联锁冲突、功率排行、再生制动估算"],
            ["路由与 UI 状态", "Vue Router + UI Store", "当前场景、Presentation 模式、Toast、确认弹窗、移动端侧栏状态"],
        ],
    )

    doc.add_paragraph("二、系统与连接")
    add_table(
        doc,
        ["展示项", "协议/后端字段", "前端字段", "展示位置", "来源", "必填"],
        [
            ["连接状态", "—", "connected", "GlobalStatusStrip、ConnectionBadge、SystemAlertBar、各页 header", "WS 连接状态", "是"],
            ["连接中", "—", "connecting", "GlobalStatusStrip、ConnectionBadge、SystemAlertBar", "WS 生命周期", "是"],
            ["数据过期", "—", "dataStale", "GlobalStatusStrip、ConnectionBadge、SystemAlertBar", "前端 >2s 无 tick", "是"],
            ["最近更新时刻", "—", "lastTickAt / freshnessLabel", "顶部状态条、DefaultLayout 顶栏", "前端收到 tick 时记录", "是"],
            ["系统模式", "system.status / system_mode", "systemMode", "Dashboard 顶栏、Signal 页、GlobalStatusStrip、SystemAlertBar", "WS", "是"],
            ["数据源", "system.data_source", "dataSource", "顶部状态条、DefaultLayout 顶栏", "WS", "否"],
            ["协议版本", "protocol_version", "protocolVersion", "Store 保留 / 调试", "WS", "否"],
            ["当前场景", "route.meta.title", "sceneLabel", "GlobalStatusStrip", "路由 + UI Store", "是"],
            ["连接错误", "WebSocket error", "lastError", "SystemAlertBar", "WS / 浏览器", "否"],
            ["演示模式", "—", "presentationMode", "全局全屏展示切换", "前端 UI", "否"],
            ["当前时钟", "—", "clockLabel", "DefaultLayout 顶栏", "前端每秒刷新", "否"],
        ],
    )

    doc.add_paragraph("三、列车 trains[] -> vehicles[]")
    doc.add_paragraph("parking_phase 枚举：cruising / approaching / braking / docking / stopped")
    add_table(
        doc,
        ["展示项", "协议字段", "前端字段", "单位", "展示位置", "来源", "必填"],
        [
            ["车号", "vehicle_id", "vehicle_id", "—", "所有地图、座舱、表格、排行", "WS", "是"],
            ["线路", "line_id", "line_id", "—", "协议保留", "WS", "否"],
            ["里程位置", "position", "position", "m", "地图、图表、表格、MA 计算", "WS", "是"],
            ["速度", "speed", "speed", "km/h", "速度表、图表、表格、制动曲线", "WS", "是"],
            ["加速度", "acceleration", "acceleration", "m/s²", "表格、Power 页估算、再生制动估算", "WS", "否"],
            ["驾驶模式", "mode", "mode", "manual / ato / atp", "座舱、车辆页、控车状态", "WS", "是"],
            ["运行中", "is_running", "is_running", "bool", "协议保留", "WS", "否"],
            ["紧急制动", "emergency_brake", "emergency_brake", "bool", "地图红车、EB 标签、事件时间线、Power 页", "WS", "是"],
            ["MA 终点", "ma_limit", "ma_limit", "m", "地图 MA 线、车辆页、Signal 页 MA 表 / 条", "WS", "是"],
            ["目标限速", "target_speed", "target_speed", "km/h", "速度表、目标速度对比、制动曲线", "WS", "是"],
            ["能耗", "energy_kwh", "energy_kwh", "kWh", "EnergyTable、Power 页负荷卡片", "WS", "否"],
            ["停车距离", "stop_distance", "stop_distance", "m", "停车距离条、制动曲线估算", "WS", "是"],
            ["下一站 / 站台", "station_name", "station_name", "—", "座舱、停车记录", "WS", "是"],
            ["进站阶段", "parking_phase", "parking_phase", "枚举", "ParkingPhaseBar", "WS（无则前端推导）", "是"],
            ["对标误差", "stop_error_cm", "stop_error_cm", "cm", "ParkingPrecision", "WS（无则前端估算）", "是"],
            ["站台 ID", "platform_id", "platform_id", "—", "协议保留", "WS", "否"],
            ["更新时间", "updated_at", "updated_at", "秒", "协议保留", "WS", "否"],
            ["选中车", "—", "selectedVehicleId / selectedVehicle", "—", "跨页面联动摘要、地图选中态", "前端 UI", "是"],
            ["车辆颜色", "—", "vehicleColor(id)", "—", "地图 / 图表 / 标签配色", "前端固定调色板", "是"],
        ],
    )

    doc.add_paragraph("四、闭塞分区 sections[] -> track_segments[]")
    add_table(
        doc,
        ["展示项", "协议字段", "前端字段", "展示位置", "来源", "必填"],
        [
            ["区段编号", "section_id", "segment_id", "区段表、Track 柱状图、Signal 冲突列表", "静态 + WS 合并", "是"],
            ["起点里程", "start", "start", "地图染色、里程条、冲突描述", "静态", "是"],
            ["终点里程", "end", "end", "同上", "静态", "是"],
            ["是否占用", "occupied", "occupied", "红 / 黄 / 绿染色、Track 表、Signal 冲突检查", "WS", "是"],
            ["显示色", "aspect", "aspect", "地图、站序图、SegmentStatusBar、Track 筛选", "WS（无则从 occupied 推）", "是"],
            ["区段状态", "condition", "condition", "协议保留", "WS", "否"],
            ["占用车辆", "vehicle_id", "occupied_by", "Track 表、Signal 冲突说明", "WS", "否"],
            ["Seg 拓扑 ID", "track_seg_id", "track_seg_id", "电子地图 Seg 定位、道岔关联", "仅静态 blocks[]", "是"],
        ],
    )

    doc.add_paragraph("五、信号机 signals[]")
    add_table(
        doc,
        ["展示项", "协议字段", "前端字段", "展示位置", "来源", "必填"],
        [
            ["信号机编号", "signal_id", "signal_id", "电子地图、Signal 页 grid、冲突列表", "静态 + WS 合并", "是"],
            ["里程位置", "position", "position", "地图定位、Signal 页、冲突邻近判断", "静态 layout", "是"],
            ["灯色", "state", "state", "三灯显示 red / yellow / green", "WS 动态", "是"],
            ["信号类型", "signal_type", "signal_type", "静态保留 / 后续 tooltip", "仅静态", "否"],
        ],
    )

    doc.add_paragraph("六、道岔 switches[] -> turnouts[]")
    add_table(
        doc,
        ["展示项", "协议/静态字段", "前端字段", "展示位置", "来源", "必填"],
        [
            ["道岔编号", "switch_id / turnout_id", "turnout_id", "地图菱形、道岔表、冲突列表", "静态 + WS 合并", "是"],
            ["定 / 反位", "routing / state", "state", "地图、表格、联锁冲突判断", "WS", "是"],
            ["锁闭", "locked", "locked", "地图颜色、表格、冲突判断", "WS", "是"],
            ["关联区段", "related_section", "related_section", "协议保留", "WS", "否"],
            ["公里标", "position", "position", "里程沙盘、冲突邻近判断", "仅静态", "是"],
            ["拓扑坐标", "graph_x / graph_y", "graph_x / graph_y", "电子地图", "仅静态", "是"],
            ["合流区段", "merge_seg_id", "merge_seg_id", "Signal 页联锁冲突检查", "仅静态", "是"],
            ["定位分支", "normal_seg", "normal_seg", "Signal 页联锁冲突检查", "仅静态", "是"],
            ["反位分支", "reverse_seg", "reverse_seg", "Signal 页联锁冲突检查", "仅静态", "是"],
        ],
    )

    doc.add_paragraph("七、供电 power")
    add_table(
        doc,
        ["展示项", "协议字段", "前端字段", "单位", "展示位置", "来源", "必填"],
        [
            ["接触网电压", "voltage", "power.voltage", "V", "Dashboard 卡片、座舱电压表、Power 页", "WS", "是"],
            ["总电流", "current", "power.current", "A", "Dashboard 卡片、Power 页", "WS", "是"],
            ["牵引功率", "power", "power.power", "kW", "Dashboard 卡片、Power 页", "WS", "是"],
            ["变电所", "substation_id", "power.substation_id", "—", "协议保留", "WS", "否"],
            ["供电故障", "is_fault", "power.is_fault", "bool", "电压表、Power 页、SystemAlertBar、事件时间线", "WS", "是"],
        ],
    )

    doc.add_paragraph("八、告警 alarms[]")
    add_table(
        doc,
        ["展示项", "协议字段", "前端字段", "展示位置", "来源", "必填"],
        [
            ["告警 ID", "alarm_id", "alarm_id", "AlarmList、事件时间线去重键", "WS", "是"],
            ["级别", "level", "level", "AlarmList 样式、事件时间线级别", "WS", "是"],
            ["来源", "source", "source", "AlarmList、事件时间线", "WS", "是"],
            ["关联车", "vehicle_id", "vehicle_id", "AlarmList、事件时间线", "WS", "否"],
            ["内容", "message", "message", "AlarmList、事件时间线标题", "WS", "是"],
            ["时间", "timestamp", "timestamp", "AlarmList、事件时间线", "WS", "是"],
        ],
    )

    doc.add_paragraph("九、静态线路 line-layout.json（地图几何）")
    add_table(
        doc,
        ["展示项", "JSON 字段", "展示位置", "来源", "必填"],
        [
            ["全线长度", "total_length_m", "各页里程比例、Track 总览", "静态", "是"],
            ["车站列表", "stations[]：station_id / name / position / graph_x / graph_y", "站序图、地图站名、Track 车站表、按站筛选", "静态", "是"],
            ["闭塞分区几何", "blocks[]：segment_id / track_seg_id / start / end", "区段合并、地图 Seg 映射、Track 表", "静态", "是"],
            ["信号静态信息", "signals[]：signal_id / position / signal_type", "地图定位、Signal 页、联锁邻近判断", "静态", "是"],
            ["道岔静态关系", "turnouts[]：turnout_id / position / graph_x / graph_y / merge_seg_id / normal_seg / reverse_seg", "地图、Signal 联锁冲突检查", "静态", "是"],
            ["Seg 拓扑", "graph.edges[]：seg_id / x1 / y1 / x2 / y2 / branch", "电子地图轨线", "静态", "是"],
            ["坡度", "slope_profile[]：position / slope", "Track 页坡度图", "静态", "否"],
            ["站台", "platforms[]", "协议保留", "静态", "否"],
        ],
    )

    doc.add_paragraph("十、前端自算 / 历史（非协议字段）")
    add_table(
        doc,
        ["展示项", "前端字段", "展示位置", "计算方式"],
        [
            ["速度历史", "vehicleHistory[id].speed", "Dashboard / Vehicle 趋势图", "每 tick 记录"],
            ["位置历史", "vehicleHistory[id].position", "Dashboard / Vehicle 趋势图", "每 tick 记录"],
            ["电压历史", "voltageHistory[]", "Dashboard / Power VoltageChart", "每 tick 记录"],
            ["时间标签", "timeLabels[]", "所有趋势图", "每 tick 格式化"],
            ["占线时间轴", "occupancyHistory[]", "Dashboard / Track OccupancyTimeline", "48 bin × 列车位置"],
            ["停车记录", "parkingRecords[]", "VehicleCockpit、Vehicle ParkingPrecision", "停稳且误差 <=50cm 时记录"],
            ["事件时间线", "eventTimeline[]", "Dashboard EventTimeline", "模式切换 / 供电故障 / 告警开闭 / EB / MA 收缩"],
            ["控车反馈", "lastControlCommand", "Dashboard 座舱、Vehicle 最近反馈", "POST control 请求结果"],
            ["控制指令历史", "controlHistory[]", "Vehicle 最近控制指令历史", "每次控车记录 pending / ok / error"],
            ["供电故障前后快照", "powerTransitions[]", "Power 页表格 / 柱状图", "power.is_fault 变化时记录前后电压电流"],
            ["当前对标误差", "currentStopErrorCm", "座舱 / Vehicle ParkingPrecision", "无 stop_error_cm 时由 stop_distance 估算"],
            ["速度监督值", "speedDelta / speedUsage", "Vehicle 目标速度 vs 实际速度", "speed 与 target_speed 计算"],
            ["MA 裕量", "maRemaining / maUsage / maRemainingText", "Vehicle / Signal 摘要", "ma_limit 与 position 计算"],
            ["制动曲线", "brakingCurveOption", "Vehicle ECharts", "根据当前速度、停车距离估算减速度并采样"],
            ["联锁冲突列表", "interlockingConflicts[]", "Signal 页", "道岔静态分支 + 区段占用 + 邻近信号组合判断"],
            ["轨道站间映射", "stationRanges / segmentStationName()", "Track 页所属区域 / 按站查看", "根据静态 station.position 划分区间"],
            ["区段筛选结果", "filteredSegments / filteredStations", "Track 页", "按站 + 异常状态 + 关键词组合过滤"],
            ["分车功率估算", "estimatePower() / vehiclePowerRanking[]", "Power 页负荷排行 / 占比图", "速度与加速度估算"],
            ["再生制动估算", "estimateRegenerativePower() / regenerativeVehicles[] / regenRatioPercent", "Power 页回馈占比", "减速度与速度估算"],
            ["顶部状态文案", "freshnessLabel / clockLabel", "DefaultLayout 顶栏、GlobalStatusStrip", "lastTickAt 与当前时钟计算"],
            ["操作反馈 UI", "toasts[] / confirmDialog", "全局 Toast / ConfirmActionDialog", "UI Store 记录"],
        ],
    )

    doc.add_paragraph("十一、按页面汇总")
    add_table(
        doc,
        ["页面", "主要展示数据"],
        [
            ["Dashboard", "三地图 + 受控车数字座舱 + 供电卡片 + 告警 + 事件时间线 + 占线时间轴 + 速度 / 位置 / 电压图 + EnergyTable"],
            ["Track", "三地图 + 区段筛选（按站 / 状态 / 关键词）+ 区段明细表 + 车站表 + 坡度图 + 区段柱状图 + 占线时间轴"],
            ["Signal", "系统模式 + 闭塞条 + 当前关注列车摘要 + 联锁冲突检查 + 信号机 grid + 道岔表 + MA 表 / 条"],
            ["Power", "供电四指标 + 当前关注列车摘要 + 再生制动回馈占比 + 分车功率排行 + 电压表 + 电压趋势 + 多车牵引负荷估算 + 功率占比图 + 故障前后电压对比"],
            ["Vehicle", "车辆选择 + 选中车详情 + 手动控车操作 + 目标速度 vs 实际速度 + MA 裕量 + 最近控制指令历史 + 制动曲线 + 停车精度 + 全车状态表 + 速度 / 位置图"],
            ["全局", "GlobalStatusStrip（连接 / 更新时间 / 当前场景 / 系统模式 / 数据源）+ SystemAlertBar + Toast + ConfirmActionDialog + Presentation 模式"],
        ],
    )

    doc.add_paragraph("十二、各组最小供给清单（联调用）")
    add_table(
        doc,
        ["组别", "必须提供", "对应前端展示"],
        [
            ["C 网关", "dashboard_snapshot 整包 WS 推送 + 稳定心跳", "几乎全部实时数据"],
            ["A 列车", "trains[] 中 vehicle_id / position / speed / acceleration / mode / ma_limit / target_speed / parking_phase / stop_error_cm 等", "地图、座舱、车辆页、Signal MA、Power 估算"],
            ["B 信号", "sections[].aspect / occupied、signals[].state、switches[].routing / locked", "Track 染色、Signal 页、联锁冲突检查"],
            ["D 前端静态", "line-layout.json 中 stations / blocks / signals / turnouts / graph / slope_profile", "地图几何、车站、分支关系、按站筛选"],
            ["供电组", "power 对象 + is_fault", "Dashboard / Power 页"],
            ["控车接口", "POST /api/v1/vehicle/control", "Vehicle 按钮、键盘控车、控制历史、Toast 反馈"],
            ["系统信息", "system.status / system.data_source / protocol_version", "顶部状态条、顶栏、告警栏"],
        ],
    )

    return doc


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: update_frontend_data_doc.py <docx-path>")
        return 1

    target = Path(sys.argv[1])
    if not target.exists():
        print(f"File not found: {target}")
        return 1

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = target.with_name(f"{target.stem}_备份_{stamp}{target.suffix}")
    shutil.copy2(target, backup)

    doc = build_document()
    doc.save(target)

    print(f"Updated: {target}")
    print(f"Backup: {backup}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
